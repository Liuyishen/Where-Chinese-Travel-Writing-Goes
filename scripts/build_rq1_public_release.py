#!/usr/bin/env python3
"""Export disclosure-safe RQ1 aggregates from private research workbooks.

This script publishes only counts, normalized entity names and entity-level
metrics. It deliberately omits full text, quotations, evidence locations,
source URLs, titles and author names.
"""
import argparse
import csv
import json
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import load_workbook


def rows(sheet):
    iterator = sheet.iter_rows(values_only=True)
    headers = next(iterator)
    for values in iterator:
        if any(value not in (None, "") for value in values):
            yield dict(zip(headers, values))


def sheet_by_prefix(workbook, prefix):
    """Use the current canonical sheet without hard-coding a row-count suffix."""
    matches = [name for name in workbook.sheetnames if name.startswith(prefix)]
    if len(matches) != 1:
        raise ValueError(f"Expected exactly one sheet beginning {prefix!r}, found {matches!r}")
    return workbook[matches[0]]


def write_csv(path, fieldnames, records):
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames, lineterminator="\n")
        writer.writeheader()
        writer.writerows(records)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--fulltext", required=True, type=Path)
    parser.add_argument("--metadata", required=True, type=Path)
    parser.add_argument("--places", required=True, type=Path)
    parser.add_argument("--atomics", required=True, type=Path)
    parser.add_argument("--output-dir", required=True, type=Path)
    parser.add_argument("--release-date", default="2026-10-03")
    args = parser.parse_args()
    args.output_dir.mkdir(parents=True, exist_ok=True)

    fulltext_book = load_workbook(args.fulltext, read_only=True, data_only=True)
    relation_rows = list(rows(sheet_by_prefix(fulltext_book, "规范关系_")))

    by_era = defaultdict(lambda: {"text_ids": set(), "relations": 0, "states": Counter()})
    for row in relation_rows:
        era = row["年代"] or "未定"
        bucket = by_era[era]
        bucket["text_ids"].add(row.get("canonical_text_id") or row["text_id"])
        bucket["relations"] += 1
        bucket["states"][row["地点状态"]] += 1

    era_order = ["1980s", "1990s", "2000s", "2010s", "2020s", "未定"]
    fulltext_eras = []
    for era in era_order:
        bucket = by_era.get(era)
        if not bucket:
            continue
        fulltext_eras.append({
            "decade": era,
            "text_count": len(bucket["text_ids"]),
            "text_place_relations": bucket["relations"],
            "visited": bucket["states"]["visited"],
            "mentioned": bucket["states"]["mentioned"],
            "imagined_or_compared": bucket["states"]["imagined_or_compared"],
            "uncertain": bucket["states"]["uncertain"],
        })
    write_csv(
        args.output_dir / "rq1_fulltext_by_decade.csv",
        list(fulltext_eras[0].keys()),
        fulltext_eras,
    )

    place_book = load_workbook(args.places, read_only=True, data_only=True)
    entity_rows = list(rows(sheet_by_prefix(place_book, "地点主表_")))
    repeated = []
    for row in entity_rows:
        if (row["关联文本数"] or 0) >= 2:
            repeated.append({
                "place_id": row["place_id"],
                "normalized_place": row["规范地点名"],
                "country_or_region": row["国家地区"],
                "text_count": row["关联文本数"],
                "author_count": row["独立作者数"],
                "visited_text_count": row["visited文本数"],
            })
    repeated.sort(key=lambda row: (-row["text_count"], -row["author_count"], row["normalized_place"] or ""))
    write_csv(
        args.output_dir / "rq1_repeated_places.csv",
        ["place_id", "normalized_place", "country_or_region", "text_count", "author_count", "visited_text_count"],
        repeated,
    )

    place_rows = entity_rows
    map_points = [{
        "place_id": row["place_id"],
        "normalized_place": row["规范地点名"],
        "country_or_region": row["国家地区"],
        "entity_type": row["实体类型"],
        "longitude": row["经度"],
        "latitude": row["纬度"],
        "text_count": row["关联文本数"],
        "author_count": row["独立作者数"],
        "visited_text_count": row["visited文本数"],
        "geometry_status": "osm_candidate_pending_human_review",
    } for row in place_rows if row["地图准入"] == "point_layer_yes"]
    map_points.sort(key=lambda row: (-int(row["text_count"]), row["normalized_place"] or ""))
    write_csv(
        args.output_dir / "rq1_candidate_map_points.csv",
        ["place_id", "normalized_place", "country_or_region", "entity_type", "longitude", "latitude", "text_count", "author_count", "visited_text_count", "geometry_status"],
        map_points,
    )

    metadata_book = load_workbook(args.metadata, read_only=True, data_only=True)
    metadata_era_rows = list(rows(metadata_book["年代证据结构"]))
    metadata_eras = [{
        "decade": row["年代"],
        "metadata_place_signals": row["地点信号行数"],
        "distinct_records": row["涉及独立记录数"],
        "specific_place_or_route_signals": row["其中具体地点或路线"],
        "regional_signals": row["其中区域覆盖线索"],
        "retrieval_only_signals": row["仅追索线索"],
    } for row in metadata_era_rows]
    write_csv(
        args.output_dir / "rq1_bibliographic_signals_by_decade.csv",
        list(metadata_eras[0].keys()),
        metadata_eras,
    )

    atomic_book = load_workbook(args.atomics, read_only=True, data_only=True)
    atomic_rows = list(rows(atomic_book["原子组件结构"]))
    atomic_summary = [{
        "component_kind": row["组件类别"],
        "entity_readiness": row["逐实体就绪度"],
        "component_count": row["原子组件数"],
        "interpretation_boundary": "Rule-based split only; not a normalized place entity or an author-visit relation.",
    } for row in atomic_rows]
    write_csv(
        args.output_dir / "rq1_bibliographic_atomic_components.csv",
        ["component_kind", "entity_readiness", "component_count", "interpretation_boundary"],
        atomic_summary,
    )

    summary = {
        "release_date": args.release_date,
        "fulltext_layer": {
            "verified_texts": len({row.get("canonical_text_id") or row["text_id"] for row in relation_rows}),
            "text_place_relations": len(relation_rows),
            "normalized_place_entities": len(entity_rows),
            "repeated_place_entities": len(repeated),
            "candidate_map_points": len(map_points),
            "rights_boundary": "Counts derive from manually coded research records; no full texts are released. Internal analysis permission, where present, is not a public redistribution licence.",
        },
        "bibliographic_layer": {
            "metadata_place_signals": sum(row["metadata_place_signals"] for row in metadata_eras),
            "specific_place_or_route_signals": sum(row["specific_place_or_route_signals"] for row in metadata_eras),
            "atomic_components": sum(int(row["component_count"]) for row in atomic_summary),
            "rights_boundary": "Bibliographic signals are metadata_only and uncertain. They are not author-visit points and are never combined with full-text relation counts.",
        },
    }
    with (args.output_dir / "rq1_release_summary.json").open("w", encoding="utf-8") as handle:
        json.dump(summary, handle, ensure_ascii=False, indent=2)
        handle.write("\n")


if __name__ == "__main__":
    main()
