#!/usr/bin/env python3
"""Build disclosure-safe, signed Q2 theme--narration aggregates.

The script joins text-level human-confirmed theme/narrative codes to the
separately signed text-place relation layer.  It writes only aggregate files.
"""

from __future__ import annotations

import argparse
import csv
import json
from collections import Counter, defaultdict
from pathlib import Path

import matplotlib.pyplot as plt
from matplotlib import font_manager
from openpyxl import load_workbook


ERAS = ["1980s", "1990s", "2000s", "2010s", "2020s"]
THEMES = [
    "自然景观", "城市日常", "移动基础设施", "人物相遇", "历史遗产", "战争/政治",
    "劳动/生计", "消费/旅游业", "边疆/民族", "生态/环境", "宗教/仪式", "自我/情感",
]
NARRATIVES = ["移动行程", "现场观察", "人物相遇", "历史文化解释", "新闻/田野报道", "回忆/停留", "mixed"]
MODES = ["road", "walking", "air", "water", "rail", "fieldwork_or_reporting", "migration_or_residence", "border_crossing", "running", "unspecified"]
PLACE_STATES = ["visited", "mentioned", "imagined_or_compared", "uncertain"]


def sheet_rows(workbook, sheet_name, header_row=1):
    sheet = workbook[sheet_name]
    iterator = sheet.iter_rows(min_row=header_row, values_only=True)
    headers = next(iterator)
    return [dict(zip(headers, row)) for row in iterator if any(value not in (None, "") for value in row)]


def split_themes(value):
    return [item.strip() for item in str(value or "").split("、") if item.strip()]


def write_csv(path, rows, headers):
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=headers, lineterminator="\n")
        writer.writeheader()
        writer.writerows(rows)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--relations-workbook", required=True, type=Path)
    parser.add_argument("--theme-workbook", required=True, type=Path)
    parser.add_argument("--output-dir", required=True, type=Path)
    parser.add_argument("--figure-dir", required=True, type=Path)
    args = parser.parse_args()
    args.output_dir.mkdir(parents=True, exist_ok=True)
    args.figure_dir.mkdir(parents=True, exist_ok=True)

    relation_book = load_workbook(args.relations_workbook, read_only=True, data_only=True)
    relations = sheet_rows(relation_book, "规范关系_410")
    reviews = {row["canonical_relationship_id"]: row for row in sheet_rows(relation_book, "关系AI复核与人工签署_20261006")}
    retained = [row for row in relations if reviews[row["canonical_relationship_id"]]["人工签署决定"] == "保留"]
    if len(retained) != 479:
        raise ValueError(f"Expected 479 retained relation rows, got {len(retained)}")

    theme_book = load_workbook(args.theme_workbook, read_only=True, data_only=True)
    coding_rows = sheet_rows(theme_book, "文本编码队列", header_row=5)
    text_codes = {}
    text_eras = {}
    for row in coding_rows:
        text_id = row["canonical_text_id"]
        themes = split_themes(row["人工确认主题"])
        narrative = str(row["人工确认叙事方式"] or "").strip()
        if row["人工签署状态"] not in {"保留", "修订"} or not themes or not narrative:
            raise ValueError(f"Incomplete signed coding for {text_id}")
        if len(themes) > 3 or any(theme not in THEMES for theme in themes) or narrative not in NARRATIVES:
            raise ValueError(f"Invalid signed coding for {text_id}: {themes}, {narrative}")
        text_codes[text_id] = {"themes": themes, "narrative": narrative}
        text_eras[text_id] = row["年代"]
    relation_text_ids = {row["canonical_text_id"] for row in retained}
    if not relation_text_ids.issubset(text_codes):
        raise ValueError("A retained-relation text is missing signed theme coding")

    theme_text_counts = Counter(theme for code in text_codes.values() for theme in code["themes"])
    narrative_text_counts = Counter(code["narrative"] for code in text_codes.values())
    theme_relation_counts = Counter()
    narrative_relation_counts = Counter()
    theme_mode_counts = Counter()
    theme_place_state_counts = Counter()
    for relation in retained:
        code = text_codes[relation["canonical_text_id"]]
        narrative_relation_counts[code["narrative"]] += 1
        mode = relation["移动方式"] or "unspecified"
        place_state = relation["地点状态"] or "unspecified"
        for theme in code["themes"]:
            theme_relation_counts[theme] += 1
            theme_mode_counts[(theme, mode)] += 1
            theme_place_state_counts[(theme, place_state)] += 1

    summary_rows = []
    for theme in THEMES:
        summary_rows.append({
            "section": "theme_texts", "category": theme, "count": theme_text_counts[theme], "denominator": len(text_codes),
            "share": round(theme_text_counts[theme] / len(text_codes), 6),
            "definition": "Texts carrying this signed theme; texts may carry one to three themes.",
        })
        summary_rows.append({
            "section": "theme_relation_support", "category": theme, "count": theme_relation_counts[theme], "denominator": len(retained),
            "share": round(theme_relation_counts[theme] / len(retained), 6),
            "definition": "Retained text-place relations attached to texts carrying this signed theme; not mutually exclusive across themes.",
        })
    for narrative in NARRATIVES:
        summary_rows.append({
            "section": "narrative_texts", "category": narrative, "count": narrative_text_counts[narrative], "denominator": len(text_codes),
            "share": round(narrative_text_counts[narrative] / len(text_codes), 6),
            "definition": "Texts with this signed primary narrative function; one function per text.",
        })
        summary_rows.append({
            "section": "narrative_relation_support", "category": narrative, "count": narrative_relation_counts[narrative], "denominator": len(retained),
            "share": round(narrative_relation_counts[narrative] / len(retained), 6),
            "definition": "Retained text-place relations carried by texts with this signed narrative function.",
        })
    write_csv(args.output_dir / "rq2_theme_narration_summary.csv", summary_rows,
              ["section", "category", "count", "denominator", "share", "definition"])

    era_rows = []
    for era in ERAS:
        era_texts = [text_id for text_id, text_era in text_eras.items() if text_era == era]
        for theme in THEMES:
            text_count = sum(theme in text_codes[text_id]["themes"] for text_id in era_texts)
            relation_count = sum(row["年代"] == era and theme in text_codes[row["canonical_text_id"]]["themes"] for row in retained)
            era_rows.append({"era": era, "theme": theme, "text_count": text_count, "signed_relation_support": relation_count,
                             "share_of_era_texts": round(text_count / len(era_texts), 6) if era_texts else 0})
    write_csv(args.output_dir / "rq2_theme_by_era.csv", era_rows,
              ["era", "theme", "text_count", "signed_relation_support", "share_of_era_texts"])

    mobility_rows = []
    for theme in THEMES:
        for mode in MODES:
            mobility_rows.append({"theme": theme, "mobility_mode": mode, "relation_count": theme_mode_counts[(theme, mode)],
                                  "definition": "A retained relation contributes once to each signed theme on its text; theme columns are not mutually exclusive."})
    write_csv(args.output_dir / "rq2_theme_mobility_matrix.csv", mobility_rows,
              ["theme", "mobility_mode", "relation_count", "definition"])

    place_rows = []
    observed_place_states = [place_state for place_state in PLACE_STATES if any(theme_place_state_counts[(theme, place_state)] for theme in THEMES)]
    for theme in THEMES:
        for place_state in observed_place_states:
            place_rows.append({"theme": theme, "place_state": place_state, "relation_count": theme_place_state_counts[(theme, place_state)],
                               "definition": "A retained relation contributes once to each signed theme on its text; theme columns are not mutually exclusive."})
    write_csv(args.output_dir / "rq2_theme_place_state_matrix.csv", place_rows,
              ["theme", "place_state", "relation_count", "definition"])

    chinese_font = "/System/Library/Fonts/Hiragino Sans GB.ttc"
    if Path(chinese_font).exists():
        font_manager.fontManager.addfont(chinese_font)
        font_family = font_manager.FontProperties(fname=chinese_font).get_name()
    else:
        font_family = "DejaVu Sans"
    plt.rcParams.update({"font.family": font_family, "font.size": 10})
    plot_themes = sorted(THEMES, key=lambda theme: (theme_text_counts[theme], theme), reverse=True)
    figure, axis = plt.subplots(figsize=(7.8, 5.8))
    bars = axis.barh(plot_themes[::-1], [theme_text_counts[theme] for theme in plot_themes[::-1]], color="#2D6C91")
    axis.set_title("Signed themes across 107 archive texts", loc="left", weight="bold")
    axis.set_xlabel("Texts carrying theme (multi-label)")
    for bar in bars:
        axis.text(bar.get_width() + 0.7, bar.get_y() + bar.get_height() / 2, str(int(bar.get_width())), va="center")
    axis.set_xlim(0, max(theme_text_counts.values()) * 1.16)
    axis.spines[["top", "right"]].set_visible(False)
    figure.text(0.98, 0.015, "107 signed texts · themes are not mutually exclusive", ha="right", va="bottom", color="#4E4A52", fontsize=9)
    figure.tight_layout(rect=[0, 0.035, 1, 1])
    figure.savefig(args.figure_dir / "q2_signed_theme_texts_20261025.svg", format="svg")
    plt.close(figure)

    visible_modes = [mode for mode in MODES if any(theme_mode_counts[(theme, mode)] for theme in THEMES)]
    matrix = [[theme_mode_counts[(theme, mode)] for mode in visible_modes] for theme in plot_themes]
    figure, axis = plt.subplots(figsize=(9.2, 6.1))
    image = axis.imshow(matrix, cmap="Purples", aspect="auto")
    axis.set_title("Mobility modes attached to signed theme-bearing texts", loc="left", weight="bold")
    axis.set_yticks(range(len(plot_themes)), plot_themes)
    axis.set_xticks(range(len(visible_modes)), visible_modes, rotation=35, ha="right")
    for y, row in enumerate(matrix):
        for x, value in enumerate(row):
            if value:
                axis.text(x, y, str(value), ha="center", va="center", fontsize=8, color="#1F2937")
    figure.colorbar(image, ax=axis, label="Retained relation support")
    figure.text(0.5, 0.015, "Theme overlaps: relation support is intentionally non-exclusive", ha="center", va="bottom", color="#4E4A52", fontsize=9)
    figure.tight_layout(rect=[0, 0.035, 1, 1])
    figure.savefig(args.figure_dir / "q2_theme_mobility_20261025.svg", format="svg")
    plt.close(figure)

    summary = {
        "release_date": "2026-10-25",
        "signed_texts": len(text_codes),
        "signed_retained_relations": len(retained),
        "signed_theme_labels": sum(theme_text_counts.values()),
        "theme_text_counts": dict(theme_text_counts),
        "narrative_text_counts": dict(narrative_text_counts),
        "interpretation_boundary": "Themes are text-level codes. Relation support joins every retained text-place relation to each signed theme on its parent text, so theme totals are intentionally non-exclusive and do not identify causal effects, trips, or population trends.",
    }
    (args.output_dir / "rq2_theme_narration_summary.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(summary, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
