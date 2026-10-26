#!/usr/bin/env python3
"""Build disclosure-safe RQ3 carrier profile from signed text-level coding."""

from __future__ import annotations

import argparse
import csv
import json
from collections import Counter
from pathlib import Path

import matplotlib.pyplot as plt
from matplotlib import font_manager
from openpyxl import load_workbook


ERAS = ["1980s", "1990s", "2000s", "2010s", "2020s"]
CARRIERS = ["报纸原刊", "报纸网页档案", "文学期刊", "杂志", "图书章节", "数字新闻平台"]
COLORS = {"报纸原刊": "#2D6C91", "报纸网页档案": "#669BBC", "文学期刊": "#8E5A99", "杂志": "#B78B40", "图书章节": "#5B4B8A", "数字新闻平台": "#2F6B5F"}


def sheet_rows(workbook, sheet_name, header_row=1):
    sheet = workbook[sheet_name]
    iterator = sheet.iter_rows(min_row=header_row, values_only=True)
    headers = next(iterator)
    return [dict(zip(headers, row)) for row in iterator if any(value not in (None, "") for value in row)]


def write_csv(path, rows, headers):
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=headers, lineterminator="\n")
        writer.writeheader()
        writer.writerows(rows)


def tidy_svg(path):
    path.write_text("\n".join(line.rstrip() for line in path.read_text(encoding="utf-8").splitlines()) + "\n", encoding="utf-8")


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--relations-workbook", required=True, type=Path)
    parser.add_argument("--carrier-workbook", required=True, type=Path)
    parser.add_argument("--output-dir", required=True, type=Path)
    parser.add_argument("--figure-dir", required=True, type=Path)
    args = parser.parse_args()
    args.output_dir.mkdir(parents=True, exist_ok=True)
    args.figure_dir.mkdir(parents=True, exist_ok=True)

    relation_book = load_workbook(args.relations_workbook, read_only=True, data_only=True)
    texts = sheet_rows(relation_book, "规范文本_107")
    relations = sheet_rows(relation_book, "规范关系_410")
    reviews = {row["canonical_relationship_id"]: row for row in sheet_rows(relation_book, "关系AI复核与人工签署_20261006")}
    retained = [row for row in relations if reviews[row["canonical_relationship_id"]]["人工签署决定"] == "保留"]

    carrier_book = load_workbook(args.carrier_workbook, read_only=True, data_only=True)
    carrier_rows = sheet_rows(carrier_book, "现有107篇载体编码", header_row=5)
    coding = {}
    for row in carrier_rows:
        text_id = row["canonical_text_id"]
        if row["人工签署状态"] not in {"保留", "修订"}:
            raise ValueError(f"Unsigned carrier row: {text_id}")
        coding[text_id] = {"carrier": row["人工确认载体类型"], "status": row["原载确认状态"]}
    text_ids = {row["canonical_text_id"] for row in texts}
    if set(coding) != text_ids:
        raise ValueError("Carrier coding does not match canonical text set")

    coverage_rows = []
    distribution_rows = []
    relation_rows = []
    for era in ERAS:
        era_texts = [row for row in texts if row["年代"] == era]
        strict = [row for row in era_texts if coding[row["canonical_text_id"]]["status"] == "可进入严格媒介分母"]
        strict_ids = {row["canonical_text_id"] for row in strict}
        coverage_rows.append({
            "era": era,
            "canonical_texts": len(era_texts),
            "strict_carrier_texts": len(strict),
            "unresolved_or_reprint_texts": len(era_texts) - len(strict),
            "strict_carrier_coverage": round(len(strict) / len(era_texts), 6),
            "definition": "Only texts with a signed, confirmed original-carrier type enter the strict carrier denominator.",
        })
        for carrier in CARRIERS:
            text_count = sum(coding[row["canonical_text_id"]]["carrier"] == carrier for row in strict)
            relation_count = sum(row["年代"] == era and row["canonical_text_id"] in strict_ids and coding[row["canonical_text_id"]]["carrier"] == carrier for row in retained)
            distribution_rows.append({"era": era, "carrier_type": carrier, "text_count": text_count, "share_of_strict_carrier_texts": round(text_count / len(strict), 6) if strict else 0,
                                      "definition": "Distribution among signed strict-carrier texts only; not a publishing-market share."})
            relation_rows.append({"era": era, "carrier_type": carrier, "retained_relation_support": relation_count,
                                  "definition": "Retained text-place relations attached to signed strict-carrier texts; not trips or independent observations."})
    write_csv(args.output_dir / "rq3_carrier_coverage_by_era.csv", coverage_rows,
              ["era", "canonical_texts", "strict_carrier_texts", "unresolved_or_reprint_texts", "strict_carrier_coverage", "definition"])
    write_csv(args.output_dir / "rq3_carrier_distribution_by_era.csv", distribution_rows,
              ["era", "carrier_type", "text_count", "share_of_strict_carrier_texts", "definition"])
    write_csv(args.output_dir / "rq3_carrier_relation_support_by_era.csv", relation_rows,
              ["era", "carrier_type", "retained_relation_support", "definition"])

    total_strict = sum(row["原载确认状态"] == "可进入严格媒介分母" for row in carrier_rows)
    total_counts = Counter(row["人工确认载体类型"] for row in carrier_rows if row["原载确认状态"] == "可进入严格媒介分母")
    summary = {
        "release_date": "2026-10-26",
        "canonical_texts": len(texts),
        "strict_carrier_texts": total_strict,
        "unresolved_or_reprint_texts": len(texts) - total_strict,
        "carrier_text_counts": dict(total_counts),
        "interpretation_boundary": "Carrier results describe the signed, confirmed-carrier subset only. Unresolved original carriers and reprint pages are excluded, so these results cannot estimate the media market or the full historical distribution of Chinese travel writing.",
    }
    (args.output_dir / "rq3_carrier_summary.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")

    chinese_font = "/System/Library/Fonts/Hiragino Sans GB.ttc"
    if Path(chinese_font).exists():
        font_manager.fontManager.addfont(chinese_font)
        font_family = font_manager.FontProperties(fname=chinese_font).get_name()
    else:
        font_family = "DejaVu Sans"
    plt.rcParams.update({"font.family": font_family, "font.size": 10})
    figure, axis = plt.subplots(figsize=(10.0, 5.0))
    x = list(range(len(ERAS)))
    bottoms = [0] * len(ERAS)
    for carrier in CARRIERS:
        values = [next(row["text_count"] for row in distribution_rows if row["era"] == era and row["carrier_type"] == carrier) for era in ERAS]
        axis.bar(x, values, bottom=bottoms, label=carrier, color=COLORS[carrier], width=0.68)
        bottoms = [bottom + value for bottom, value in zip(bottoms, values)]
    for index, row in enumerate(coverage_rows):
        axis.text(index, bottoms[index] + 0.4, f"n={row['strict_carrier_texts']}", ha="center", va="bottom", fontsize=9)
    axis.set_ylim(0, max(bottoms) + 2.2)
    axis.set_xticks(x, ERAS)
    axis.set_ylabel("Texts with confirmed original-carrier type")
    axis.set_title("Carrier types in the signed confirmed-carrier subset", loc="left", weight="bold")
    axis.legend(frameon=False, loc="upper left", bbox_to_anchor=(1.02, 1))
    axis.spines[["top", "right"]].set_visible(False)
    figure.text(0.5, 0.015, "38 texts are unresolved or reprints and excluded from this strict comparison", ha="center", va="bottom", color="#4E4A52", fontsize=9)
    figure.tight_layout(rect=[0, 0.035, 0.78, 1])
    figure_path = args.figure_dir / "rq3_carrier_types_20261026.svg"
    figure.savefig(figure_path, format="svg")
    tidy_svg(figure_path)
    plt.close(figure)
    print(json.dumps(summary, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
