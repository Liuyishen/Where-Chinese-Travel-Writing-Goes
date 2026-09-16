#!/usr/bin/env python3
"""Create a metadata-only public CSV from a private master workbook."""
import argparse
import csv
from pathlib import Path
from openpyxl import load_workbook

FIELDS = [
    "record_id", "text_id", "source_id", "stat_year", "stat_decade", "title",
    "author", "parent_source", "carrier", "genre", "destination_or_route",
    "geographic_scope", "travel_strength", "verification_grade", "access_status",
    "rights_status", "source_url",
]
SOURCE = {
    "record_id": "记录ID", "text_id": "文本实体ID", "source_id": "母来源ID",
    "stat_year": "保守统计年", "stat_decade": "保守统计年代", "title": "篇目标题",
    "author": "作者", "parent_source": "母来源", "carrier": "出版／刊载信息",
    "genre": "写作类型_标准", "destination_or_route": "目的地／线路",
    "geographic_scope": "地域范围_标准", "travel_strength": "旅行性_标准",
    "verification_grade": "核验等级", "source_url": "来源URL",
}

def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", required=True, type=Path)
    parser.add_argument("--output", required=True, type=Path)
    parser.add_argument("--sheet", default="标准化候选池")
    args = parser.parse_args()
    workbook = load_workbook(args.input, read_only=True, data_only=True)
    sheet = workbook[args.sheet]
    headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
    positions = {name: index for index, name in enumerate(headers)}
    args.output.parent.mkdir(parents=True, exist_ok=True)
    with args.output.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=FIELDS)
        writer.writeheader()
        for row in sheet.iter_rows(min_row=2, values_only=True):
            if not any(value not in (None, "") for value in row):
                continue
            record = {field: row[positions[SOURCE[field]]] if SOURCE[field] in positions else "" for field in SOURCE}
            marker = str(row[positions.get("质量标记", -1)] or "") if "质量标记" in positions else ""
            record["access_status"] = "quote_only" if "quote_only" in marker else "metadata_only"
            record["rights_status"] = "not_verified" if "not_verified" in marker else "not_stated"
            writer.writerow(record)

if __name__ == "__main__":
    main()
